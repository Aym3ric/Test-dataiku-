import base64
import io
from docx import Document
from dataiku.llm.python import BaseLLM

DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

class DirectFileArtifactAgent(BaseLLM):
    def __init__(self):
        super().__init__()

    def process(self, query, settings, trace):
        prompt = query["messages"][-1]["content"]
        
        # Création de l'artefact
        artifacts = [self._create_word_artifact(prompt)]
        
        return {
            "ok": True,
            "text": "Voici le document généré sous forme d'artefact :",
            "finishReason": "STOP",
            "artifacts": artifacts
        }

    def _create_word_artifact(self, topic):
        # 1. Génération du fichier Word en mémoire
        doc = Document()
        doc.add_heading(f'Document : {topic}', 0)
        doc.add_paragraph(f'Ce document a été généré automatiquement sur le sujet : {topic}')
        doc.add_heading('Contenu Principal', level=1)
        doc.add_paragraph('Voici le contenu détaillé du document.')
        
        doc_buffer = io.BytesIO()
        doc.save(doc_buffer)
        doc_buffer.seek(0)
        doc_base64 = base64.b64encode(doc_buffer.getvalue()).decode('utf-8')
        
        # 2. Schéma attendu par le sérialiseur LLM Mesh Dataiku
        return {
            "id": "word-report-1",
            "type": "FILE",
            "name": "Rapport_Synthese.docx",
            "description": "Document Word généré par l'agent",
            "parts": [
                {
                    "type": "DATA_INLINE",
                    "index": 0,
                    "mimeType": DOCX_MIMETYPE,
                    "dataBase64": doc_base64
                }
            ]
        }
